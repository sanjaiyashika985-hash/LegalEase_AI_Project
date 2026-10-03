"""Create downloadable TXT, DOCX, and PDF versions of edited document text."""
from io import BytesIO

from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from xml.sax.saxutils import escape


def to_txt(text: str) -> bytes:
    return text.encode("utf-8")


def to_docx(text: str) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    for line in text.splitlines():
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(5)
        run = para.add_run(line)
        if line.strip().isupper() and len(line.strip()) < 90:
            run.bold = True
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def to_pdf(text: str) -> bytes:
    output = BytesIO()
    pdf = SimpleDocTemplate(
        output, pagesize=A4, rightMargin=22 * mm, leftMargin=22 * mm,
        topMargin=22 * mm, bottomMargin=22 * mm,
        title="LegalEase Draft", author="LegalEase",
    )
    styles = getSampleStyleSheet()
    story = []
    for line in text.splitlines():
        if not line.strip():
            story.append(Spacer(1, 5))
        else:
            safe = escape(line)
            if line.strip().isupper() and len(line.strip()) < 90:
                story.append(Paragraph(f"<b>{safe}</b>", styles["Heading3"]))
            else:
                story.append(Paragraph(safe, styles["BodyText"]))
    pdf.build(story)
    return output.getvalue()
